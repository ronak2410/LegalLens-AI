"""
LegalLens AI - FastAPI Application Server.
Provides REST APIs for legal contract analysis, document-grounded Q&A,
side-by-side comparison, negotiation playbooks, batch due diligence, and exports.
"""

import io
import json
import os
from pathlib import Path
from typing import Optional, List, Dict, Any

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request, APIRouter
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel, Field

# Optional Word Docx support
try:
    import docx
    from docx.shared import Pt, RGBColor
except ImportError:
    docx = None

from backend.samples import SAMPLE_DOCUMENTS
from backend.legal_engine import analyze_document_text
from backend.qa_engine import answer_document_question, COMMON_SUGGESTED_QUESTIONS
from backend.comparator import compare_legal_documents
from backend.portfolio_engine import analyze_portfolio_documents
from backend.parser import extract_text_from_file, MAX_FILE_SIZE_BYTES

# Base Directories
BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = BASE_DIR / "frontend"
PUBLIC_DIR = BASE_DIR / "public"

app = FastAPI(
    title="LegalLens AI API",
    description="AI-powered legal information assistant for understanding contracts, risks, and obligations.",
    version="1.0.0"
)

api_router = APIRouter()

import time
from collections import defaultdict
from starlette.middleware.gzip import GZipMiddleware

# Rate limiting settings (In-memory sliding window)
RATE_LIMIT_WINDOW_SEC = 60.0
RATE_LIMIT_MAX_REQUESTS = 120
_IP_REQUEST_HISTORY: Dict[str, List[float]] = defaultdict(list)

# ----------------- Security Middlewares -----------------

@app.middleware("http")
async def enforce_payload_size_and_security_headers(request: Request, call_next):
    # 1. Enforce 15MB maximum request payload limit
    content_length = request.headers.get("content-length")
    if content_length and int(content_length) > MAX_FILE_SIZE_BYTES:
        return JSONResponse(
            status_code=413,
            content={"detail": "Request payload exceeds maximum allowable limit of 15MB."}
        )

    # 2. Rate limit API endpoints
    if request.url.path.startswith("/api") or request.url.path in ["/analyze", "/chat", "/compare", "/upload-batch"]:
        client_ip = request.client.host if request.client else "127.0.0.1"
        now = time.time()
        history = _IP_REQUEST_HISTORY[client_ip]
        _IP_REQUEST_HISTORY[client_ip] = [t for t in history if (now - t) < RATE_LIMIT_WINDOW_SEC]
        if len(_IP_REQUEST_HISTORY[client_ip]) >= RATE_LIMIT_MAX_REQUESTS:
            return JSONResponse(
                status_code=429,
                content={"detail": "Rate limit exceeded. Maximum 120 requests per minute allowable."}
            )
        _IP_REQUEST_HISTORY[client_ip].append(now)

    # 3. Process request
    response = await call_next(request)

    # 4. Attach strict security headers and rate limit transparency headers
    client_ip = request.client.host if request.client else "127.0.0.1"
    remaining = max(0, RATE_LIMIT_MAX_REQUESTS - len(_IP_REQUEST_HISTORY.get(client_ip, [])))
    response.headers["X-RateLimit-Limit"] = str(RATE_LIMIT_MAX_REQUESTS)
    response.headers["X-RateLimit-Remaining"] = str(remaining)

    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["X-XSS-Protection"] = "1; mode=block"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Content-Security-Policy"] = (
        "default-src 'self' https://fonts.googleapis.com https://fonts.gstatic.com; "
        "style-src 'self' 'unsafe-inline' https://fonts.googleapis.com; "
        "font-src 'self' https://fonts.gstatic.com; "
        "img-src 'self' data:; "
        "connect-src 'self'; "
        "script-src 'self' 'unsafe-inline';"
    )
    return response

import urllib.parse
from starlette.types import ASGIApp, Scope, Receive, Send

class VercelPathNormalizerMiddleware:
    """
    ASGI middleware that normalizes incoming requests when deployed on Vercel Serverless.
    Translates /api/index.py or query subpaths into proper FastAPI API paths before routing.
    """
    def __init__(self, app: ASGIApp):
        self.app = app

    async def __call__(self, scope: Scope, receive: Receive, send: Send):
        if scope["type"] == "http":
            target_path = None
            query_bytes = scope.get("query_string", b"")
            if query_bytes:
                qs_str = query_bytes.decode("utf-8", errors="ignore")
                parsed_qs = urllib.parse.parse_qs(qs_str)
                if "__vercel_subpath__" in parsed_qs:
                    sub = parsed_qs["__vercel_subpath__"][0]
                    target_path = "/api/" + sub.lstrip("/") if sub else "/api"
                    cleaned_pairs = [(k, v) for k, vals in parsed_qs.items() if k != "__vercel_subpath__" for v in vals]
                    scope["query_string"] = urllib.parse.urlencode(cleaned_pairs).encode("utf-8")

            if not target_path:
                headers = dict(scope.get("headers", []))
                matched = (
                    headers.get(b"x-matched-path")
                    or headers.get(b"x-vercel-matched-path")
                    or headers.get(b"x-forwarded-uri")
                )
                if matched:
                    decoded = matched.decode("utf-8", errors="ignore").split("?")[0]
                    if decoded.startswith("/api"):
                        target_path = decoded

            curr_path = scope.get("path", "")
            if target_path:
                scope["path"] = target_path
                if "raw_path" in scope:
                    scope["raw_path"] = target_path.encode("utf-8")
            elif curr_path.startswith("/api/index.py"):
                sub = curr_path[len("/api/index.py"):]
                norm = "/api" + (sub if sub.startswith("/") else ("/" + sub if sub else ""))
                scope["path"] = norm
                if "raw_path" in scope:
                    scope["raw_path"] = norm.encode("utf-8")

        await self.app(scope, receive, send)

# Add Vercel ASGI path normalizer
app.add_middleware(VercelPathNormalizerMiddleware)

# High-efficiency GZip compression middleware
app.add_middleware(GZipMiddleware, minimum_size=1000)

# Standard CORS policy without insecure wildcard credentials
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)



# ----------------- Request Models -----------------

class PlaybookRule(BaseModel):
    regex_pattern: str = Field(..., max_length=150)
    title: str = Field(..., max_length=100)
    severity: str = "High"
    category: str = "Custom Playbook"
    description: Optional[str] = ""
    impact: Optional[str] = ""
    recommendation: Optional[str] = ""


class AnalyzeTextRequest(BaseModel):
    text: str
    filename: Optional[str] = "Contract.txt"
    api_key: Optional[str] = None
    playbook_rules: Optional[List[PlaybookRule]] = None


class BatchDocumentItem(BaseModel):
    text: str
    filename: Optional[str] = "Contract.txt"


class BatchAnalyzeRequest(BaseModel):
    documents: List[BatchDocumentItem]
    playbook_rules: Optional[List[PlaybookRule]] = None


class ChatRequest(BaseModel):
    question: str = Field(..., max_length=500)
    document_text: str
    document_title: Optional[str] = "Uploaded Document"
    api_key: Optional[str] = None


class CompareRequest(BaseModel):
    doc1_text: str
    doc2_text: str
    doc1_name: Optional[str] = "Document A"
    doc2_name: Optional[str] = "Document B"


# ----------------- API Endpoints -----------------

@api_router.get("/health")
@api_router.get("/api/health")
def health_check():
    """Service health and metadata check."""
    return {
        "status": "healthy",
        "service": "LegalLens AI",
        "version": "1.0.0",
        "challenge": "PromptWars Virtual: AI for Legal Assistance & Access"
    }


@api_router.get("/samples")
@api_router.get("/api/samples")
def get_sample_list():
    """Returns catalog of pre-configured realistic contracts."""
    summaries = []
    for key, doc in SAMPLE_DOCUMENTS.items():
        summaries.append({
            "id": doc["id"],
            "title": doc["title"],
            "doc_type": doc["doc_type"],
            "parties": doc["parties"],
            "effective_date": doc["effective_date"],
            "summary": doc["summary"],
            "word_count": len(doc["text"].split())
        })
    return {"samples": summaries}


@api_router.get("/samples/{sample_id}")
@api_router.get("/api/samples/{sample_id}")
def get_sample_detail(sample_id: str):
    """Retrieves full text and metadata for a specific sample."""
    if sample_id not in SAMPLE_DOCUMENTS:
        raise HTTPException(status_code=404, detail="Sample contract not found.")
    return SAMPLE_DOCUMENTS[sample_id]


@api_router.post("/analyze")
@api_router.post("/api/analyze")
def analyze_contract(req: AnalyzeTextRequest):
    """Analyzes document text and returns structured legal assessment."""
    if not req.text or not req.text.strip():
        raise HTTPException(status_code=400, detail="Document text cannot be empty.")

    playbook_dicts = [rule.model_dump() for rule in req.playbook_rules] if req.playbook_rules else None
    analysis = analyze_document_text(req.text, req.filename or "Contract.txt", playbook_dicts, req.api_key)
    return analysis


@api_router.post("/analyze-upload")
@api_router.post("/api/analyze-upload")
async def analyze_uploaded_file(file: UploadFile = File(...), playbook_rules: Optional[str] = Form(None), api_key: Optional[str] = Form(None)):
    """Receives uploaded PDF, DOCX, or TXT file and analyzes contents."""
    filename = file.filename or "Uploaded Document"
    content_bytes = await file.read()

    try:
        text = extract_text_from_file(content_bytes, filename)
    except ValueError as val_err:
        raise HTTPException(status_code=413, detail=str(val_err))

    if len(text.strip()) < 50:
        raise HTTPException(status_code=400, detail="Could not extract sufficient text from the uploaded file. Please paste text or upload a readable PDF/DOCX/TXT.")

    playbook_dicts = None
    if playbook_rules:
        try:
            playbook_dicts = json.loads(playbook_rules)
        except Exception:
            pass

    analysis = analyze_document_text(text, filename, playbook_dicts, api_key)
    analysis["extracted_text"] = text
    return analysis


@api_router.post("/analyze-batch")
@api_router.post("/api/analyze-batch")
def analyze_contract_batch(req: BatchAnalyzeRequest):
    """Batch analyzes multiple contracts concurrently for Portfolio Due Diligence."""
    if not req.documents or len(req.documents) == 0:
        raise HTTPException(status_code=400, detail="At least one document is required for batch analysis.")

    playbook_dicts = [rule.model_dump() for rule in req.playbook_rules] if req.playbook_rules else None
    doc_items = [{"text": d.text, "filename": d.filename} for d in req.documents if d.text.strip()]
    if not doc_items:
        raise HTTPException(status_code=400, detail="Provided documents contain no readable text.")

    portfolio_results = analyze_portfolio_documents(doc_items, playbook_dicts)
    return portfolio_results


@api_router.post("/upload-batch")
@api_router.post("/api/upload-batch")
async def upload_contract_batch(files: List[UploadFile] = File(...), playbook_rules: Optional[str] = Form(None)):
    """Receives multiple uploaded files (PDF/DOCX/TXT) and performs Portfolio Due Diligence."""
    if not files:
        raise HTTPException(status_code=400, detail="No files uploaded.")

    playbook_dicts = None
    if playbook_rules:
        try:
            playbook_dicts = json.loads(playbook_rules)
        except Exception:
            pass

    doc_items = []
    for file in files:
        filename = file.filename or "Contract"
        content_bytes = await file.read()
        try:
            text = extract_text_from_file(content_bytes, filename)
            if text and len(text.strip()) >= 50:
                doc_items.append({"text": text, "filename": filename})
        except ValueError:
            continue

    if not doc_items:
        raise HTTPException(status_code=400, detail="Could not extract readable text from any uploaded files.")

    portfolio_results = analyze_portfolio_documents(doc_items, playbook_dicts)
    return portfolio_results


@api_router.post("/chat")
@api_router.post("/api/chat")
def ask_document(req: ChatRequest):
    """Grounded Q&A strictly against document text."""
    if not req.document_text.strip():
        raise HTTPException(status_code=400, detail="Active document text is missing.")

    response = answer_document_question(req.question, req.document_text, req.document_title or "Uploaded Contract", req.api_key)
    return response


@api_router.post("/compare")
@api_router.post("/api/compare")
def compare_contracts(req: CompareRequest):
    """Side-by-side comparison between two documents."""
    if not req.doc1_text.strip() or not req.doc2_text.strip():
        raise HTTPException(status_code=400, detail="Both documents are required for comparison.")

    comparison = compare_legal_documents(
        req.doc1_text,
        req.doc2_text,
        req.doc1_name or "Document A",
        req.doc2_name or "Document B"
    )
    return comparison


@api_router.get("/suggested-questions")
@api_router.get("/api/suggested-questions")
def get_suggested_questions():
    """Returns categorized suggested questions for legal document exploration."""
    return {"categories": COMMON_SUGGESTED_QUESTIONS}


@api_router.post("/export-markdown")
@api_router.post("/api/export-markdown")
@api_router.post("/export-report")
@api_router.post("/api/export-report")
def export_legal_report_markdown(req: AnalyzeTextRequest):
    """Generates a downloadable Markdown legal briefing report."""
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Document text is required.")

    playbook_dicts = [rule.model_dump() for rule in req.playbook_rules] if req.playbook_rules else None
    analysis = analyze_document_text(req.text, req.filename or "Contract.txt", playbook_dicts)
    meta = analysis["metadata"]
    risk = analysis["risk_assessment"]

    report_lines = [
        f"# LegalLens AI Briefing Report: {meta['title']}",
        f"**Date Generated:** {meta.get('effective_date', 'Current Session')} | **Document Type:** {meta['doc_type']}",
        f"**Overall Risk Profile:** {risk['level']} ({risk['score']}/100)",
        "\n> **LEGAL DISCLAIMER:** LegalLens AI provides general legal information and is not a law firm or substitute for qualified legal advice.\n",
        "## Executive Plain-Language Summary",
    ]
    for s in analysis["plain_language_summary"]:
        report_lines.append(f"### {s['title']}")
        report_lines.append(f"{s['simple_explanation']}")
        report_lines.append(f"**Action:** {s['what_to_do']}\n")

    report_lines.append("## Identified Risks & Contractual Asymmetries")
    for r in analysis["risks_and_flags"]:
        report_lines.append(f"- **[{r['severity']} Risk] {r['title']}** ({r['clause_ref']})")
        report_lines.append(f"  - *Impact:* {r['impact']}")
        report_lines.append(f"  - *Negotiation Tip:* {r['recommendation']}")

    report_lines.append("\n## Proposed Counter-Language / Redlines")
    for red in analysis.get("redlines", []):
        report_lines.append(f"### Proposed Redline: {red['title']}")
        report_lines.append(f"```text\n{red['proposed_clause']}\n```")

    report_lines.append("\n## Questions for Your Attorney Consultation")
    for i, q in enumerate(analysis["lawyer_questions"], 1):
        report_lines.append(f"{i}. {q}")

    report_lines.append("\n## Pre-Signing Actionable Checklist")
    for chk in analysis["checklist"]:
        report_lines.append(f"- [ ] **[{chk['priority']}]** {chk['label']}")

    return {"markdown": "\n".join(report_lines), "filename": f"{meta['title'].replace(' ', '_')}_LegalLens_Report.md"}


@api_router.post("/export-docx")
@api_router.post("/api/export-docx")
def export_legal_report_docx(req: AnalyzeTextRequest):
    """Generates a downloadable DOCX legal briefing report with inline redlines."""
    if not req.text.strip():
        raise HTTPException(status_code=400, detail="Document text is required.")

    playbook_dicts = [rule.model_dump() for rule in req.playbook_rules] if req.playbook_rules else None
    analysis = analyze_document_text(req.text, req.filename or "Contract.txt", playbook_dicts)
    meta = analysis["metadata"]
    risk = analysis["risk_assessment"]

    if not docx:
        md_res = export_legal_report_markdown(req)
        f = io.BytesIO(md_res["markdown"].encode('utf-8'))
        headers = {'Content-Disposition': f'attachment; filename="{meta["title"].replace(" ", "_")}_Report.txt"'}
        return StreamingResponse(f, media_type="text/plain", headers=headers)

    doc = docx.Document()
    doc.add_heading(f"LegalLens AI Briefing: {meta['title']}", 0)
    doc.add_paragraph(f"Date Generated: {meta.get('effective_date', 'Current Session')} | Document Type: {meta['doc_type']}")

    risk_p = doc.add_paragraph()
    risk_p.add_run("Overall Risk Profile: ").bold = True
    risk_run = risk_p.add_run(f"{risk['level']} ({risk['score']}/100)")
    risk_run.font.color.rgb = RGBColor(0xef, 0x44, 0x44) if risk['level'] == 'High' else RGBColor(0x10, 0xb9, 0x81)

    doc.add_heading("Executive Plain-Language Summary", level=1)
    for s in analysis["plain_language_summary"]:
        doc.add_heading(s['title'], level=2)
        doc.add_paragraph(s['simple_explanation'])
        p = doc.add_paragraph()
        p.add_run("Action: ").bold = True
        p.add_run(s['what_to_do'])

    doc.add_heading("Identified Risks & Contractual Asymmetries", level=1)
    for r in analysis["risks_and_flags"]:
        doc.add_heading(f"[{r['severity']} Risk] {r['title']} ({r['clause_ref']})", level=2)
        doc.add_paragraph(f"Impact: {r['impact']}")
        doc.add_paragraph(f"Negotiation Tip: {r['recommendation']}")

    doc.add_heading("Proposed Counter-Language / Redlines", level=1)
    for red in analysis.get("redlines", []):
        doc.add_heading(f"Proposed Redline: {red['title']}", level=2)
        p = doc.add_paragraph()
        run = p.add_run(red['proposed_clause'])
        run.font.color.rgb = RGBColor(0x10, 0xb9, 0x81)

    doc.add_heading("Pre-Signing Actionable Checklist", level=1)
    for chk in analysis["checklist"]:
        doc.add_paragraph(f"[ ] [{chk['priority']}] {chk['label']}", style='List Bullet')

    doc.add_paragraph("LEGAL DISCLAIMER: LegalLens AI provides general legal information and is not a law firm or substitute for qualified legal advice.", style='Intense Quote')

    f = io.BytesIO()
    doc.save(f)
    f.seek(0)
    headers = {'Content-Disposition': f'attachment; filename="{meta["title"].replace(" ", "_")}_LegalLens_Report.docx"'}
    return StreamingResponse(f, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers)


# Include API router
app.include_router(api_router)


# ----------------- Frontend Safe Static Files Mount -----------------
STATIC_DIR = PUBLIC_DIR if PUBLIC_DIR.exists() else FRONTEND_DIR
if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

    @app.get("/")
    async def serve_index():
        return FileResponse(str(STATIC_DIR / "index.html"), media_type="text/html")

    @app.get("/static/{file_path:path}")
    async def serve_static_direct(file_path: str):
        target = (STATIC_DIR / file_path).resolve()
        if target.is_relative_to(STATIC_DIR.resolve()) and target.exists() and target.is_file():
            media_type = None
            if file_path.endswith(".css"):
                media_type = "text/css"
            elif file_path.endswith(".js"):
                media_type = "application/javascript"
            elif file_path.endswith(".png"):
                media_type = "image/png"
            elif file_path.endswith(".svg"):
                media_type = "image/svg+xml"
            return FileResponse(str(target), media_type=media_type)
        raise HTTPException(status_code=404, detail="Static asset not found")

    @app.api_route("/{catchall:path}", methods=["GET", "POST", "OPTIONS"])
    async def serve_frontend(request: Request, catchall: str):
        # Do NOT serve index.html for API paths
        if catchall.startswith("api/") or catchall == "api" or catchall.startswith("analyze") or catchall.startswith("chat") or catchall.startswith("compare") or catchall.startswith("upload") or catchall.startswith("export"):
            raise HTTPException(status_code=404, detail="API route not found")


        if request.method != "GET":
            raise HTTPException(status_code=405, detail="Method Not Allowed")

        try:
            cleaned = catchall.replace("static/", "", 1) if catchall.startswith("static/") else catchall
            target_path = (STATIC_DIR / cleaned).resolve()
            if target_path.is_relative_to(STATIC_DIR.resolve()) and target_path.exists() and target_path.is_file():
                media_type = None
                if cleaned.endswith(".css"):
                    media_type = "text/css"
                elif cleaned.endswith(".js"):
                    media_type = "application/javascript"
                elif cleaned.endswith(".png"):
                    media_type = "image/png"
                elif cleaned.endswith(".svg"):
                    media_type = "image/svg+xml"
                return FileResponse(str(target_path), media_type=media_type)
        except Exception:
            pass
        return FileResponse(str(STATIC_DIR / "index.html"), media_type="text/html")


